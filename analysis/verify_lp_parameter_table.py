"""Verify preservation before delivering the narrowly scoped LP workbook edit."""
from copy import copy
import hashlib
import json
from pathlib import Path
import shutil
import sys

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SUPPORT = ROOT / ('outputs/20261013_lp_confirmed/support' if '--confirmed' in sys.argv else 'outputs/20261013_lp_review/support')


def main():
    payload = json.loads((SUPPORT / 'rows.json').read_text(encoding='utf-8'))
    output = SUPPORT.parent / '全部原始编码.xlsx'
    original = openpyxl.load_workbook(payload['template'])
    updated = openpyxl.load_workbook(output)
    s, t = original.worksheets[0], updated.worksheets[0]
    rows = [list(r) for r in t.iter_rows(min_row=2, max_row=len(payload['rows']) + 1, max_col=6, values_only=True)]
    assert rows == payload['rows'] and rows[0] == payload['example']
    assert all(r[0] and r[1] for r in rows)
    assert len({r[1] for r in rows}) == len(rows)
    assert all(c.value is None for row in t.iter_rows(min_row=len(rows) + 2) for c in row)
    assert original.sheetnames == updated.sheetnames and s.sheet_format == t.sheet_format
    assert {k: d.width for k, d in s.column_dimensions.items()} == {k: d.width for k, d in t.column_dimensions.items()}
    assert {k: d.height for k, d in s.row_dimensions.items()} == {k: d.height for k, d in t.row_dimensions.items()}
    assert not any(c.alignment.wrap_text for row in t for c in row)
    assert s.freeze_panes == t.freeze_panes
    assert list(s.tables) == list(t.tables)
    for name in s.tables:
        assert s.tables[name].ref == t.tables[name].ref
    for i in (1, 2):
        assert list(original.worksheets[i].values) == list(updated.worksheets[i].values)
    for r in range(1, len(payload['original_rows']) + 2):
        for c in range(1, 7):
            assert all(copy(getattr(s.cell(r, c), n)) == copy(getattr(t.cell(r, c), n))
                       for n in ('font', 'fill', 'border', 'alignment', 'number_format'))
    others = lambda values: [r for r in values if r[4] not in {'hrsg1_lp', 'hrsg2_lp'}
                             and r[1] not in payload.get('authorized_new_interface_fields', [])]
    assert others(payload['original_rows']) == others(rows)
    for path, digest in payload['source_point_hashes'].items():
        assert hashlib.sha256((ROOT / path).read_bytes()).hexdigest() == digest
    source = Path(payload['source'])
    out_hash = hashlib.sha256(output.read_bytes()).hexdigest()
    src_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    assert src_hash in {payload['source_sha256'], out_hash}, 'Current workbook changed; reconcile before delivery'
    report = {**payload['counts'], 'example_and_other_systems_preserved': True,
              'all_spacing_styles_and_blank_sheets_preserved': True, 'source_point_workbooks_unchanged': True,
              'backup': payload['template'], 'output_sha256': out_hash}
    if '--deliver' in sys.argv:
        shutil.copy2(output, source)
        assert source.read_bytes() == output.read_bytes()
        report['delivered_to'] = str(source)
    elif src_hash == out_hash:
        report['delivered_to'] = str(source)
    original.close()
    updated.close()
    (ROOT / 'analysis/lp_parameter_table_verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False))


if __name__ == '__main__':
    main()
